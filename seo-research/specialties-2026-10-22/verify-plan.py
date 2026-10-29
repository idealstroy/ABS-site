from pathlib import Path
from copy import copy
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook
import json

root=Path(__file__).resolve().parents[2]
qa=root/'tmp/specialties-2026-10-22'
data=json.loads((Path(__file__).parent/'research.json').read_text(encoding='utf-8'))
core_path=root/'The structure and semantics of the website/Семантическое ядро Сайт АБС.xlsx'
core=load_workbook(core_path)
before=load_workbook(qa/('before-'+core_path.name))
assert core.sheetnames[:len(before.sheetnames)]==before.sheetnames
for sheet in before:
    for row in sheet:
        for c in row:
            n=core[sheet.title][c.coordinate]
            assert c.value==n.value,(sheet.title,c.coordinate)
            if c.value is not None:
                assert copy(c.font)==copy(n.font)
                assert copy(c.fill)==copy(n.fill)
                assert copy(c.alignment)==copy(n.alignment)
                assert c.number_format==n.number_format
cached=load_workbook(core_path,data_only=True)
assert cached['машинист автовышки и автогидроп']['D29'].value==626
for page in data['pages']:
    s=core[page['sheet']]
    assert s.max_row==2+len(page['observed'])+len(page['masks'])
    for r in range(3,s.max_row+1):
        assert s.cell(r,3).value is None and s.cell(r,4).value is None
        assert s.cell(r,2).value.date().isoformat()=='2026-10-22'
        assert s.cell(r,10).value=='Не проверялась'
structure=load_workbook(root/'The structure and semantics of the website/struktura_saita_ABS_SEO_updated.xlsx')
for page in data['pages']:
    for name in ('Структура сайта','Карта сайта','Сводка кластеров','Паспорт структуры','Интенты и границы','Перелинковка','Региональный шаблон'):
        assert any(page['path']==c.value for row in structure[name] for c in row),(page['path'],name)
assert structure['Контроль пункта 1']['B5'].value=="=COUNTA('Паспорт структуры'!A5:A22)"
d=Document(root/'правила/Дорожная карта создания сайта АБС.docx')
paragraphs=[p.text for p in d.paragraphs]
for page in data['pages']:
    assert paragraphs.count('ПУНКТ '+page['step']+' [ЗАПЛАНИРОВАН]')==1
    assert any(page['path'] in p for p in paragraphs)
    assert any(page['step']==row.cells[0].text for row in d.tables[2].rows)
    assert any(page['step']==row.cells[0].text for row in d.tables[-1].rows)
for f in (core_path,root/'The structure and semantics of the website/struktura_saita_ABS_SEO_updated.xlsx',root/'правила/Дорожная карта создания сайта АБС.docx'):
    with ZipFile(f) as z:assert z.testzip() is None
print(json.dumps({'pages_planned':len(data['pages']),'observed_and_named_phrases':sum(len(p['observed'])for p in data['pages']),'separate_masks':sum(len(p['masks'])for p in data['pages']),'existing_core_preserved':True,'roadmap_steps':['18а','18б'],'docx_visual_qa':'unavailable: bundled LibreOffice missing'},ensure_ascii=False))
