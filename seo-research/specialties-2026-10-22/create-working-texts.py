from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

folder=Path(__file__).resolve().parent
root=folder.parents[1]
data=json.loads((folder/'content-drafts.json').read_text(encoding='utf-8'))
outputs=[]
for page in data['pages']:
    target=root/'Textes'/('Текст страницы - '+page['name']+'.docx')
    if target.exists():raise RuntimeError('Working DOCX already exists; edit it in place instead of regenerating: '+str(target))
    doc=Document()
    section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(20)
    section.left_margin=section.right_margin=Mm(22)
    for name,size in [('Normal',11),('Title',20),('Subtitle',12),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
        s=doc.styles[name]
        s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(7)
        s.paragraph_format.line_spacing=1.12
        if name.startswith('Heading'):s.paragraph_format.keep_with_next=True
    doc.add_paragraph(page['name'],style='Title')
    doc.add_paragraph('Текст страницы',style='Subtitle')
    doc.add_paragraph('Редакция от 22 октября 2026 года')
    doc.add_heading('SEO параметры',level=1)
    for label,value in [('URL','/mashinisty-spetstehniki/'+page['slug']+'/'),('Title',page['seo']['title']),('Description',page['seo']['description'])]:
        p=doc.add_paragraph();p.add_run(label+': ').bold=True;p.add_run(value)
    doc.add_heading('Первый экран',level=1)
    doc.add_heading(page['h1'],level=2)
    doc.add_paragraph(page['intro'])
    doc.add_paragraph('Кнопка: '+page['cta'])
    for b in page['blocks']:
        doc.add_heading(b['title'],level=1)
        if b.get('intro'):doc.add_paragraph(b['intro'])
        for item in b.get('items',[]):
            doc.add_heading(item['title'],level=2)
            doc.add_paragraph(item['text'])
        for line in b.get('bullets',[]):doc.add_paragraph(line,style='List Bullet')
        if b.get('note'):doc.add_paragraph(b['note'])
        if b.get('form'):
            doc.add_paragraph('Поля: Имя, Телефон, Регион, Комментарий.')
            doc.add_paragraph('Отправляя форму, вы соглашаетесь на обработку персональных данных.')
            doc.add_paragraph('Кнопка: Оставить заявку.')
    doc.add_heading('Частые вопросы',level=1)
    for f in page['faq']:
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run(f['question']).bold=True
        doc.add_paragraph(f['answer'])
    doc.add_heading('Внутренние ссылки',level=1)
    for link in page['links']:doc.add_paragraph(link['label']+' — '+link['url'])
    doc.core_properties.title=page['name']
    doc.core_properties.subject='Рабочий текст страницы для согласования'
    doc.core_properties.author='АБС ИНЖИНИРИНГ'
    doc.save(target)
    check=Document(target)
    assert sum(p.text=='Оставьте заявку' for p in check.paragraphs)==1
    assert len([f for f in page['faq'] if any(p.text==f['answer']for p in check.paragraphs)])==5
    outputs.append(str(target))
print(json.dumps({'created':outputs},ensure_ascii=False))
