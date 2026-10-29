from pathlib import Path
import json
from docx import Document

folder=Path(__file__).resolve().parent
root=folder.parents[1]
registry=root/'src/data/planned-specialties.json'
data=json.loads(registry.read_text(encoding='utf-8'))
for page in data['pages']:
    page['status']='awaiting_text_approval'
    page['frequencyStatus']='Bukvarix проверен 22.10.2026, Весь мир. Прямые B2B-маски не найдены; измеренные профессиональные фразы и исключения добавлены в ядро.'
    page['workingText']='Textes/Текст страницы - '+page['title']+'.docx'
    page['seoBrief']='seo-research/specialties-2026-10-22/seo-brief.md'
registry.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
research=folder/'research.json'
data=json.loads(research.read_text(encoding='utf-8'))
data['verification']={'date':'2026-10-22','source':'Bukvarix','region':'Весь мир','cache_date':None,'raw_rows':154,'queries':15,'zero_result_masks':13,'raw':'raw/','records':'verified-keywords.json','log':'bukvarix-query-log.json','note':'Отсутствие ключей в базе не означает нулевую частотность. Общие профессиональные запросы не подтверждают B2B-спрос.'}
research.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
path=root/'правила/Дорожная карта создания сайта АБС.docx'
doc=Document(path)
active=False
for p in doc.paragraphs:
    if p.text.startswith('ПУНКТ 18а') or p.text.startswith('ПУНКТ 18б'):
        active=True
        for r in p.runs:r.text=r.text.replace('[ЗАПЛАНИРОВАН]','[ТЕКСТ НА СОГЛАСОВАНИИ]')
    elif p.text.startswith('ПУНКТ '):active=False
    if active and p.text.startswith('Действия: Изучить вкладку'):
        for i,r in enumerate(p.runs):
            r.text=('Действия: Изучить обновлённую вкладку специальности в книге «Семантическое ядро Сайт АБС.xlsx». Bukvarix проверен 22.10.2026, регион «Весь мир». Профессиональные фразы, вакансии и справочные запросы разделены. По прямым B2B-маскам ключи не найдены; числовые нули не присвоены. Перед SEO-решениями учитывать отсутствие измерения по России.'if i==0 else '')
for row in doc.tables[-1].rows:
    if row.cells[0].text in ('18а','18б'):
        row.cells[2].text='Текст на согласовании'
        row.cells[5].text='Bukvarix проверен. SEO-ТЗ и рабочий DOCX подготовлены. Частотности — Весь мир. Вёрстка после согласования текста.'
doc.save(path)
print('Progress updated: two text drafts await approval')
