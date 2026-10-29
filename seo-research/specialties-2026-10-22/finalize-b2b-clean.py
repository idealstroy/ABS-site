from pathlib import Path
from copy import copy
from zipfile import ZipFile
import csv, hashlib, json, shutil
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[2]
folder=Path(__file__).resolve().parent
qa=root/'tmp/specialties-b2b-clean-2026-10-22'
name='Семантическое ядро Сайт АБС.xlsx'
target=root/'The structure and semantics of the website'/name
backup=qa/('before-'+name)
merged=qa/('merged-'+name)
report=json.loads((qa/'clean-report.json').read_text(encoding='utf-8'))
before=load_workbook(backup)
after=load_workbook(merged)
assert before.sheetnames==after.sheetnames
for s in before:
    n=after[s.title]
    if s.title in report['rowLimits']:
        assert n.max_row==report['rowLimits'][s.title]
        assert list(n.tables.values())[0].ref==f'A2:M{n.max_row}'
        expected=[r['keyword'] for r in report['kept'] if r['values'][11] in [n['L3'].value]]
        assert [n.cell(i,1).value for i in range(3,n.max_row+1)]==expected
        for i in range(3,n.max_row+1):
            assert n.cell(i,2).number_format==s['B3'].number_format
        continue
    assert s.max_row==n.max_row and s.max_column==n.max_column
    for row in s:
        for c in row:
            nc=n[c.coordinate]
            assert c.value==nc.value,(s.title,c.coordinate)
            assert copy(c.font)==copy(nc.font)
            assert copy(c.fill)==copy(nc.fill)
            assert copy(c.alignment)==copy(nc.alignment)
            assert c.number_format==nc.number_format
    for key in s.tables:
        assert s.tables[key].to_tree().attrib==n.tables[key].to_tree().attrib
assert load_workbook(merged,data_only=True)['машинист автовышки и автогидроп']['D29'].value==626
with ZipFile(merged) as z:assert z.testzip() is None
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert digest(target) in (digest(backup),digest(merged)), 'Workbook changed after backup; do not overwrite user edits'
shutil.copyfile(merged,target)
assert digest(target)==digest(merged)

fields=['page','keyword','frequency','exact_frequency','region','source','source_url','frequency_status','decision']
def compact(r):
    v=r['values']
    return dict(zip(fields,[r['page'],r['keyword'],v[2],v[3],v[4],v[5],v[6],v[9],r.get('reason',v[12])]))
for filename,rows in [('05_topic_candidates.csv',report['kept']),('06_b2b_core.csv',report['kept']),('07_removed_noncommercial.csv',report['removed'])]:
    with (folder/filename).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(compact(r) for r in rows)
research=json.loads((folder/'research.json').read_text(encoding='utf-8'))
research['verification'].update(b2bOnly=True,active_rows=15,removed_rows=154,active_measured_rows=1,active_unconfirmed_rows=14,active_core='06_b2b_core.csv',removed_archive='07_removed_noncommercial.csv',note='Рабочее ядро содержит только поиск/предоставление специалистов. Сырые 154 записи сохранены вне рабочего ядра. Частотности 14 коммерческих формулировок/масок не подтверждены; одна фраза с измерением не подтверждает B2B-объём спроса. Отсутствие ключей в базе не означает нулевую частотность.')
(folder/'research.json').write_text(json.dumps(research,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'kept':len(report['kept']),'removed':len(report['removed']),'other_sheets_unchanged':True,'main_workbook_saved':True},ensure_ascii=False))
