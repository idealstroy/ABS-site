from pathlib import Path
from copy import copy
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook
import json
root=Path(__file__).resolve().parents[2]
folder=Path(__file__).resolve().parent
qa=root/'tmp/specialties-content-2026-10-22'
filename='Семантическое ядро Сайт АБС.xlsx'
path=root/'The structure and semantics of the website'/filename
before=load_workbook(qa/('before-'+filename))
after=load_workbook(path)
assert before.sheetnames==after.sheetnames
edited={'моторист цементировочного агр','слесарь по ремонту ДСТ'}
for s in before:
    if s.title in edited:continue
    n=after[s.title]
    for row in s:
        for c in row:
            nc=n[c.coordinate]
            assert c.value==nc.value,(s.title,c.coordinate)
            if c.value is not None:
                assert copy(c.font)==copy(nc.font)
                assert copy(c.fill)==copy(nc.fill)
                assert copy(c.alignment)==copy(nc.alignment)
                assert c.number_format==nc.number_format
    for name in s.tables:
        table=s.tables[name]
        nt=n.tables[name]
        assert table.ref==nt.ref
        assert [x.name for x in table.tableColumns]==[x.name for x in nt.tableColumns]
records=json.loads((folder/'verified-keywords.json').read_text(encoding='utf-8'))
research=json.loads((folder/'research.json').read_text(encoding='utf-8'))
if research.get('verification',{}).get('b2bOnly'):
    approved=json.loads((root/'tmp/specialties-b2b-clean-2026-10-22/clean-report.json').read_text(encoding='utf-8'))['kept']
    records=[r for r in records if any(a['keyword'].lower()==r['keyword'].lower() and a['page']==r['page'] for a in approved)]
for p in research['pages']:
    s=after[p['sheet']]
    rows={str(r[0].value).lower():r for r in list(s.rows)[2:]}
    for rec in records:
        if rec['page']!=p['slug']:continue
        row=rows[rec['keyword'].lower()]
        assert row[2].value==rec['frequency']
        assert row[3].value==rec['exact_frequency']
        assert row[4].value=='Весь мир'
    table=list(s.tables.values())[0]
    assert table.ref.endswith(str(s.max_row)),(table.ref,s.max_row)
cache=load_workbook(path,data_only=True)
assert cache['машинист автовышки и автогидроп']['D29'].value==626
drafts=json.loads((folder/'content-drafts.json').read_text(encoding='utf-8'))
for p in drafts['pages']:
    file=root/'Textes'/('Текст страницы - '+p['name']+'.docx')
    doc=Document(file)
    texts=[x.text for x in doc.paragraphs]
    assert p['h1'] in texts and p['intro'] in texts
    assert texts.count('Оставьте заявку')==1
    for faq in p['faq']:assert faq['answer'] in texts
    with ZipFile(file)as z:assert z.testzip() is None
    assert not(root/'src/data/pages'/(p['slug']+'.json')).exists()
roadmap=Document(root/'правила/Дорожная карта создания сайта АБС.docx')
assert sum(p.text.endswith('[ТЕКСТ НА СОГЛАСОВАНИИ]')for p in roadmap.paragraphs)==2
print(json.dumps({'measured_keywords_verified':len(records),'old_sheets_formulas_styles_and_tables_preserved':True,'working_texts':2,'form_headings_each':1,'production_pages_created':0,'status':'awaiting_text_approval','docx_visual_qa':'not performed: bundled LibreOffice unavailable'},ensure_ascii=False))
