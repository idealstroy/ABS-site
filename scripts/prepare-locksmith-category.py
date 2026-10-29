from pathlib import Path
from datetime import date
import csv, json, re, shutil
from urllib.parse import quote
from docx import Document
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
out=root/'seo-research/slesari-category-2026-10-29'
(out/'raw').mkdir(parents=True,exist_ok=True)
seeds=['слесарь по ремонту спецтехники','слесарь ремонтник спецтехники']
synonyms=['слесарь по ремонту строительной техники','слесарь по ремонту дорожно-строительной техники','слесарь по ремонту и обслуживанию спецтехники','слесарь по обслуживанию спецтехники','автослесарь по ремонту спецтехники','автослесарь спецтехники','слесарь-механик спецтехники','слесарь-механик по ремонту спецтехники']
commercial=['аутстаффинг слесарей','аутсорсинг слесарей','предоставление слесарей','подбор слесаря','подбор слесарей','подбор персонала слесарей','отбор слесарей по ремонту спецтехники','предоставление слесарей по ремонту спецтехники','аутстаффинг слесарей по ремонту спецтехники','нанять слесаря по ремонту спецтехники','найти слесаря по ремонту спецтехники','слесари по ремонту спецтехники для компании']
checked=seeds+['автослесарь спецтехники','подбор слесаря','слесарь по ремонту строительной техники','слесарь механик спецтехники','слесарь по обслуживанию спецтехники','предоставление слесарей','аутсорсинг слесарей']
(out/'_masks_seed.txt').write_text('\n'.join(seeds+synonyms+commercial)+'\n',encoding='utf-8')
(out/'_anchors.txt').write_text('слесар\nавтослесар\nспецтехник\nстроительн\nавтомобил\n',encoding='utf-8')
raw=[]
log=[]
for mask in checked:
    src=Path('C:/Users/marke/Downloads')/(mask+'.csv')
    shutil.copy2(src,out/'raw'/src.name)
    rows=list(csv.reader(src.read_text(encoding='utf-8-sig').splitlines(),delimiter=';'))
    rows=[r for r in rows[1:] if r]
    url='https://www.bukvarix.com/keywords/?q='+quote(mask)
    log.append(dict(mask=mask,rows=len(rows),date='2026-10-29',region='Весь мир',url=url,source='Bukvarix',cache_date=None))
    for r in rows:raw.append(dict(keyword=r[0],frequency=int(r[3]),exact=int(r[4]),mask_parent=mask,source='Bukvarix',depth=0,region='Весь мир',date='2026-10-29',url=url))
log.append(dict(mask='аутстаффинг слесарей',rows=0,status='Ключевые слова не найдены; частотность неизвестна',date='2026-10-29',region='Весь мир',url='https://www.bukvarix.com/keywords/?q='+quote('аутстаффинг слесарей')))
(out/'query-log.json').write_text(json.dumps(log,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def writecsv(name,rows,fields):
    with (out/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
fields=['keyword','frequency','exact','mask_parent','source','depth','region','date','url']
writecsv('03_semantics_raw.csv',raw,fields)
seen={}
for r in raw:seen.setdefault(r['keyword'].lower().replace('ё','е'),r)
core=[];excluded=[];ambiguous=[]
for r in seen.values():
    k=r['keyword']
    if re.search('ваканс|работ[ауые]|работу|вахт|резюме|жиль|обучен|удостоверен|еткс|окз|инструкц|обязанност|разряд|фото| это',k):reason='Вакансии, обучение, справочный или нормативный интент'
    elif 'сборщик' in k:reason='Другая специальность'
    elif re.search('подбор|аутсорсинг|аутстаффинг|нанять|предоставление',k):
        if 'заявка' in k:reason='Возможный поиск образца заявки; требует отдельной проверки'
        else:core.append({**r,'target':'/slesari/','status':'Коммерческий кандидат; общие слесари, не доказательство спроса именно на спецтехнику'});continue
    else:ambiguous.append({**r,'status':'Название профессии без подтверждённого коммерческого интента'});continue
    excluded.append({**r,'reason':reason})
writecsv('04_excluded.csv',excluded,fields+['reason'])
writecsv('05_ambiguous.csv',ambiguous,fields+['status'])
writecsv('06_b2b_core.csv',core,fields+['target','status'])
writecsv('07_masks.csv',[dict(mask=m,kind='Исходная' if m in seeds else 'Синоним' if m in synonyms else 'Коммерческий модификатор',status='Проверена' if m in checked or m=='аутстаффинг слесарей' else 'Кандидат на проверку') for m in seeds+synonyms+commercial],['mask','kind','status'])

base=json.loads((root/'src/data/pages/mashinisty-spetstehniki.json').read_text(encoding='utf-8'))
base.update(output='slesari/index.html',path='/slesari/',status='draft',h1='Слесари по ремонту спецтехники и автомобилей для объектов',navigationLabel='Слесари',eyebrow='АБС — предоставление персонала',breadcrumbs=[{'label':'Главная','href':'/'},{'label':'Слесари','href':''}],schema=[],images=[],licenseSegments=[])
base['seo']=dict(title='Отбор слесарей по ремонту спецтехники и автомобилей | АБС',description='Предоставляем слесарей по ремонту спецтехники и автомобилей для объектов и ремонтных баз. Учитываем состав парка, виды ремонта, опыт, регион и график.',robots='noindex, nofollow',canonical='',openGraph=False)
base['cards']=[dict(icon='tools.svg',title='Слесарь по ремонту ДСТ',description='Для обслуживания и ремонта дорожно-строительной техники с учётом состава парка и задач ремонтной службы.',href='/_prototype/slesar-po-remontu-dst/',productionHref='/slesari/slesar-po-remontu-dst/',columnClass='col-12 col-md-6'),dict(icon='gear.svg',title='Слесарь по ремонту автомобилей',description='Для обслуживания и ремонта автомобильной техники на базе или объекте заказчика.',href='/_prototype/slesar-po-remontu-am/',productionHref='/slesari/slesar-po-remontu-am/',columnClass='col-12 col-md-6')]
c=base['category']
c.update(listTitle='Выберите специальность слесаря',heroLead='Отбираем слесарей по ремонту спецтехники и автомобилей для объектов и ремонтных баз. Уточняем состав техники, виды работ, опыт и документы специалистов, регион и график выхода.',introTitle='Слесарь под задачи ремонтной службы',intro='У слесарей одной профессии может различаться опыт работы с техникой и её узлами. До поиска согласуем, какие машины предстоит обслуживать, какие работы выполнять и насколько самостоятельно должен работать сотрудник. Предоставляем персонал для вашей ремонтной службы; перечень задач задаёте вы.',requirementsTitle='Что указать в заявке',requirements=[dict(label=s,icon=i) for s,i in [('Типы и модели техники, узлы и виды ремонта','tools.svg'),('Количество слесарей и желаемая дата выхода','people.svg'),('Регион, рабочее место и график','geo-alt.svg'),('Опыт, разряд и документы по требованиям объекта','file-earmark-check.svg')]],processTitle='Процесс отбора слесарей',process=[dict(title=t,text=s) for t,s in [('Получаем заявку','Уточняем технику, виды ремонта, количество специалистов и условия работы.'),('Согласовываем профиль','Фиксируем опыт, документы, самостоятельность, регион и график.'),('Проводим первичный отбор','Сверяем опыт кандидатов с задачами ремонтной службы и условиями объекта.'),('Представляем специалистов','Передаём сведения для рассмотрения и финального собеседования заказчика.')]],checksTitle='При отборе проверяем',checks=['Опыт обслуживания и ремонта нужных типов техники','Узлы и системы, с которыми работал кандидат','Опыт диагностики, разборки, ремонта и сборки по задачам заявки','Квалификацию и документы, предусмотренные требованиями заказчика','Готовность к региону, графику и условиям рабочего места','Возможную дату выхода на объект или ремонтную базу'],formTitle='Оставьте заявку')
truck=json.loads((root/'src/data/pages/voditel-gruzovogo-avtomobilya.json').read_text(encoding='utf-8'))['specialist']
c.update(heroImage=truck['heroImage'],heroAlt=truck['heroAlt'],heroWidth=1916,heroHeight=821,overviewImage=truck['tasksImage'],overviewImageAlt=truck['tasksImageAlt'],checksImage=truck['checksImage'],checksImageAlt=truck['checksImageAlt'])
base['faq']=[dict(question=q,answer=a) for q,a in [('Каких слесарей можно заказать?','В разделе представлены слесари по ремонту дорожно-строительной техники и автомобилей. Выберите направление и укажите машины, узлы и перечень работ.'),('Вы предоставляете слесарей или выполняете ремонт техники?','Предоставляем специалистов под задачи вашей ремонтной службы. В заявке согласуем требования к сотрудникам и условия работы на вашей базе или объекте.'),('Можно ли заказать несколько слесарей с разным опытом?','Да. Укажите количество сотрудников и распределение задач. Требования к технике, узлам, опыту и документам согласуем отдельно по каждой позиции.'),('Что нужно сообщить для начала отбора?','Укажите технику, виды ремонта, количество слесарей, регион, график и желаемую дату выхода. Отдельно добавьте требования к опыту, квалификации и документам.'),('Можно ли привлечь слесарей на удалённый объект?','До поиска уточняем место работы, оснащение, формат выездов, график и условия размещения, если они предусмотрены заявкой.')]]
copyfile=out/'category-copy.json'
if copyfile.exists():
    copy=json.loads(copyfile.read_text(encoding='utf-8'))
    base['seo'].update(copy['seo'])
    base['h1']=copy['h1']
    base['category'].update(copy['category'])
    base['faq']=copy['faq']
    for card,description in zip(base['cards'],copy['cardDescriptions']):card['description']=description
existing=root/'src/data/pages/slesari.json'
if existing.exists():
    current=json.loads(existing.read_text(encoding='utf-8'))
    if current.get('status')=='ready':
        base['status']='ready'
        base['seo'].update({k:current['seo'][k] for k in ['robots','canonical','openGraph']})
        base['schema']=current['schema']
        base['faqSchema']=current.get('faqSchema',False)
        for card in base['cards']:card['href']=card['productionHref']
(root/'src/data/pages/slesari.json').write_text(json.dumps(base,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(out/'README.md').write_text(f'''# Раздел Слесари

Дата: 29.10.2026. URL: /slesari/. Две исходные маски, повтор удалён. Всего {len(seeds+synonyms+commercial)} маски: профессиональные синонимы и коммерческие модификаторы. Варианты «слесарь-механик» и «автослесарь» рассматриваются как смежные поисковые формулировки, а не полные квалификационные эквиваленты профессии.

Источник: веб-интерфейс Bukvarix, регион в сервисе «Весь мир», дата обновления базы не показана. Регион страницы — Россия. Эти параметры не смешиваются. Получено {len(raw)} сырых строк; {len(seen)} уникальных. Коммерческие кандидаты: {len(core)}, неоднозначные названия профессий: {len(ambiguous)}, исключено: {len(excluded)}. Суммы частотностей не рассчитывались.

Исходные широкие маски содержат значительную долю вакансий и справочных запросов. В ядро не включены работа, вакансии, обучение, нормативные данные и запросы образцов заявок. Коммерческие кандидаты общего поиска слесарей не доказывают отдельный спрос именно на спецтехнику. Маски без проверки не получают выдуманную частотность. Отсутствие результата Bukvarix не означает нулевой спрос.

Исходные CSV сохранены в raw/, журнал — query-log.json. 07_masks.csv содержит статус каждой маски. Семантическая книга в этом проходе не изменена: отдельного листа категории «Слесари» в ней нет, а смешивать широкую категорию с листом ДСТ нельзя. CSV подготовлены для последующего переноса.

Родитель отвечает на общий запрос о предоставлении слесарей и помогает выбрать направление. Узкие ДСТ-запросы закреплены за /slesari/slesar-po-remontu-dst/, автомобильные — за /slesari/slesar-po-remontu-am/. Детальные тексты дочерних страниц не дублируются. При согласовании ядра дополнительно проверить коммерческий интент и целевой регион.

Страница создана как draft/noindex на действующем шаблоне категории. Карточки пока ведут на рабочие превью дочерних страниц; productionHref хранит назначенные адреса публикации. Изображения временно взяты из действующей библиотеки грузового автомобиля; уникальные изображения ремонтной службы ещё не подготовлены. Деплой не выполнялся.

Конкурентные источники для коммерческих формулировок:
- https://stroyservisgarant.ru/services/slesarno-remontnye-raboty/ — аутстаффинг слесарей-ремонтников, общее производство.
- https://rostovnadonu.leaderg.ru/autsorsing/slesarej/ — общий аутсорсинг слесарей.
- https://cat-serv.ru/servis-spectekhniki/slesar — ремонт техники как услуга, другой интент.
''',encoding='utf-8')

roadmap=root/'правила/Дорожная карта создания сайта АБС.docx'
backup=root/'tmp/before-slesari-category-roadmap.docx'
if not backup.exists():shutil.copy2(roadmap,backup)
doc=Document(roadmap)
if not any('ПУНКТ 18г' in p.text for p in doc.paragraphs):
    anchor=next(p for p in doc.paragraphs if p.text=='Раздел Слесари')
    anchor.text='ПУНКТ 18г — Создать страницу раздела Слесари'
    doc.add_paragraph('Статус: локальное превью на согласовании. URL /slesari/, уровень 2, шаблон категории. Добавлено 29 октября 2026 года по прямой команде пользователя.')
    for s in ['1. Семантика: проверить две уникальные исходные маски и дополнительные синонимы из seo-research/slesari-category-2026-10-29/07_masks.csv. Отделить коммерческий поиск специалистов от вакансий, обучения, справочных запросов и услуг ремонта. Сохранить источник, регион, дату и частотности без подмены пропусков нулями.','2. SEO-структура: общие запросы слесарей закрепить за /slesari/. ДСТ — /slesari/slesar-po-remontu-dst/, автомобили — /slesari/slesar-po-remontu-am/. Подробные тексты дочерних страниц на категорию не копировать.','3. Текст и вёрстка: использовать действующий шаблон второго уровня, две карточки специальностей, требования к заявке, процесс отбора, проверки и FAQ. Единственная встроенная форма — «Оставьте заявку».','4. Изображения: согласовать материалы по ремонту техники. В текущем превью использованы существующие изображения автомобильной техники.','5. Интеграция: после приёмки подключить раздел к главной, меню, футеру и перелинковке. Карточки направить на согласованные URL дочерних страниц. Сохранить общие формы, счётчики, верификацию и сторонние скрипты.','6. Приёмка: проверить desktop/mobile, ссылки, форму и отсутствие дублей. До согласования сохранять draft/noindex и отсутствие в sitemap. Публикация — отдельный шаг.']:
        doc.add_paragraph(s)
    table=doc.tables[-1]
    from copy import deepcopy
    from docx.table import _Row
    ref=next(r for r in table.rows if r.cells[0].text=='18в')
    tr=deepcopy(ref._tr);ref._tr.addnext(tr);row=_Row(tr,table)
    for cell,value in zip(row.cells,['18г','Создать страницу категории Слесари','Превью на согласовании','18б, 18в','Обязательна','Семантика, маски и локальная страница подготовлены. Согласовать текст, изображения и интеграцию.']):cell.text=value
    table=doc.tables[2];tr=deepcopy(table.rows[-1]._tr);table._tbl.append(tr);row=_Row(tr,table)
    for cell,value in zip(row.cells,['18г','Слесари','2','/slesari/']):cell.text=value
    doc.tables[0].cell(0,0).text='25 основных\nи 4 дополнительных пункта'
    doc.tables[0].cell(0,1).text='20\nстраниц SEO-структуры'
    doc.save(roadmap)
registry=root/'src/data/planned-specialties.json'
d=json.loads(registry.read_text(encoding='utf-8'))
for section in d.get('sections',[]):
    if section['path']=='/slesari/':section.update(status='preview',step='18г',masks=seeds+synonyms+commercial,research='seo-research/slesari-category-2026-10-29/')
registry.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(dict(raw=len(raw),unique=len(seen),core=len(core),ambiguous=len(ambiguous),excluded=len(excluded),masks=len(seeds+synonyms+commercial)),ensure_ascii=False))
