"""Map the existing AM copy to template 3 and add roadmap step 18в."""
from pathlib import Path
from copy import deepcopy
import json, shutil
from docx import Document
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
out = root / 'seo-research/slesar-am'
out.mkdir(parents=True, exist_ok=True)
source = root / 'Textes/Текст страницы - Слесарь по ремонту АМ.docx'
roadmap = root / 'правила/Дорожная карта создания сайта АБС.docx'
backup = root / 'tmp/slesar-am-before'
backup.mkdir(parents=True, exist_ok=True)
for p in (source, roadmap):
    target = backup / p.name
    if not target.exists(): shutil.copy2(p, target)
d = Document(source)
for p in d.paragraphs:
    new = p.text.replace('/slesar-po-remontu-dst/', '/slesar-po-remontu-am/').replace('Автомобилей', 'автомобилей').replace('Для каких работ отбираем слесарей ДСТ', 'Для каких работ отбираем слесарей по ремонту автомобилей')
    if new != p.text:
        p.runs[0].text = new
        for r in p.runs[1:]: r.text = ''
d.save(source)
t = [p.text for p in d.paragraphs]
data = dict(file='preview-slesar-po-remontu-am.json', slug='slesar-po-remontu-am', title='Слесарь по ремонту АМ', h1=t[8], seoTitle=t[5].removeprefix('Title: '), description=t[6].removeprefix('Description: '), heroLead=t[9], objectionsTitle=t[11], objectionsText=t[12], objections=[[t[i],t[i+1],icon] for i,icon in zip([13,15,17],['minecart-loaded.svg','people.svg','geo-alt.svg'])], checksTitle=t[19], checksText=t[20], checks=t[21:28], processTitle=t[28], process=[[t[i],t[i+1]] for i in [29,31,33,35]], resultText=t[38], resultItems=t[39:43], resultNote=t[43], formText=t[45], teamText=t[50], tasksTitle=t[51], tasksIntro=t[52], tasks=[[t[i],t[i+1],'tools.svg'] for i in [53,55,57,59]], faq=[[t[i],t[i+1]] for i in [62,64,66,68]], images=['bulldozer-specialist-hero.png','bulldozer-specialist-checks.png','bulldozer-specialist-process.png','bulldozer-specialist-tasks.png'], imageAlts=['Техника на объекте','Специалист и руководитель проверяют готовность к работе','Специалист и руководитель обсуждают задачу на объекте','Техника на площадке'])
truck=json.loads((root/'src/data/pages/voditel-gruzovogo-avtomobilya.json').read_text(encoding='utf-8'))['specialist']
data['images']=[truck[k].split('/')[-1] for k in ['heroImage','checksImage','processImage','tasksImage']]
data['imageAlts']=[truck[k] for k in ['heroAlt','checksImageAlt','processImageAlt','tasksImageAlt']]
(out / 'content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

doc = Document(roadmap)
def txt(e): return ''.join(e.itertext()) if False else ''.join(n.text or '' for n in e.iter(qn('w:t')))
def settext(e, value):
    nodes=list(e.iter(qn('w:t')))
    if nodes:
        nodes[0].text=value
        for n in nodes[1:]: n.text=''
if not any('ПУНКТ 18в' in p.text for p in doc.paragraphs):
    els=list(doc.element.body)
    start=next(i for i,e in enumerate(els) if txt(e).startswith('ПУНКТ 18б'))
    end=next(i for i,e in enumerate(els) if i>start and txt(e)=='ПУНКТ 19')
    for e in els[start:end]:
        clone=deepcopy(e)
        paras=[clone] if clone.tag==qn('w:p') else list(clone.iter(qn('w:p')))
        for p in paras:
            s=txt(p).replace('18б','18в').replace('18а','18б').replace('Слесарь по ремонту ДСТ','Слесарь по ремонту АМ').replace('слесарь по ремонту ДСТ','слесарь по ремонту автомобилей').replace('слесарей по ремонту ДСТ','слесарей по ремонту автомобилей').replace('/slesar-po-remontu-dst/','/slesar-po-remontu-am/').replace('22 октября 2026','29 октября 2026')
            if s.startswith('Действия: Изучить вкладку'): s='Действия: Проверить коммерческий кластер ремонта автомобилей в существующем ядре и отделить его от ремонта ДСТ. Частотности по АМ отдельно не проверены; не переносить показатели ДСТ.'
            if s.startswith('Действия: Использовать реестр конкурентов'): s='Действия: Для SEO-публикации отдельно проверить выдачу по предоставлению слесарей по ремонту автомобилей. Исходный текст: Textes/Текст страницы - Слесарь по ремонту АМ.docx.'
            settext(p,s)
        els[end].addprevious(clone)
    for table_index, table in enumerate([doc.tables[2],doc.tables[-1]]):
        ref=next(r for r in table.rows if r.cells[0].text=='18б')
        row=deepcopy(ref._tr);ref._tr.addnext(row)
        cells=list(row.iter(qn('w:tc')))
        values=['18в','Слесарь по ремонту АМ','3','/mashinisty-spetstehniki/slesar-po-remontu-am/'] if table_index==0 else ['18в','Создать страницу Слесарь по ремонту АМ','Превью на согласовании','18б','Обязательна','Существующий текст перенесён в шаблон 3. Локальное noindex-превью создано. Семантика и публикация требуют отдельной проверки.']
        for c,v in zip(cells,values):
            ps=list(c.iter(qn('w:p')));settext(ps[0],v)
            for p in ps[1:]:settext(p,'')
    for p in doc.paragraphs:
        s=p.text.replace('18а и 18б','18а, 18б и 18в').replace('Версия 3.1  |  Две новые специальности добавлены 22 октября 2026 года','Версия 3.2  |  Слесарь по ремонту АМ добавлен 29 октября 2026 года')
        if s!=p.text:settext(p._p,s)
    doc.tables[0].cell(0,0).text='25 основных\nи 3 дополнительных пункта'
    doc.tables[0].cell(0,1).text='19\nстраниц SEO-структуры'
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if 'Объединить 18 принятых страниц' in p.text:settext(p._p,p.text.replace('18','19'))
    doc.save(roadmap)
print('AM content and roadmap updated')
